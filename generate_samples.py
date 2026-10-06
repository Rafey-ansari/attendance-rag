"""Generate one sample attendance file per format + canonical schema + expected results.  Usage: python generate_samples.py [outdir]"""
import io, json, sys
from pathlib import Path
import pandas as pd

EMP = [("E001", "Asha Patil", "Engineering"), ("E002", "Ravi Kumar", "Engineering"), ("E003", "Meera Shah", "HR"),
       ("E004", "John Dsouza", "HR"), ("E005", "Priya Nair", "Finance")]
DAYS = {"2026-09-01": "PPPAP", "2026-09-02": "PALPP", "2026-09-03": "PPAAP", "2026-09-04": "PPPPP", "2026-09-05": "APPPP", "2026-09-08": "PPPAL"}
FULL = {"P": "Present", "A": "Absent", "L": "Leave"}
COLS = ["employee_id", "employee_name", "department", "date", "status"]
rows = lambda day: [(e[0], e[1], e[2], day, FULL[c]) for e, c in zip(EMP, DAYS[day])]

def scan_png(day):
    from PIL import Image, ImageDraw, ImageFont
    lines = ["ATTENDANCE REGISTER " + day] + [" | ".join(r) for r in rows(day)]
    try: font = ImageFont.truetype("DejaVuSans.ttf", 30)
    except OSError: font = ImageFont.load_default(size=30)
    img = Image.new("RGB", (1300, 80 + 60 * len(lines)), "white"); d = ImageDraw.Draw(img)
    for i, l in enumerate(lines): d.text((30, 30 + 60 * i), l, fill="black", font=font)
    b = io.BytesIO(); img.save(b, "PNG"); return b.getvalue(), img.size

def main(out="samples"):
    import docx, fitz
    out = Path(out); out.mkdir(exist_ok=True, parents=True)
    pd.DataFrame(rows("2026-09-01"), columns=COLS).to_csv(out / "attendance.csv", index=False)
    x = pd.DataFrame([(a, b, c, d, s[0]) for a, b, c, d, s in rows("2026-09-02")], columns=["Emp Code", "Employee", "Dept", "Date", "Attendance"])
    x.to_excel(out / "attendance.xlsx", index=False)
    d = docx.Document(); d.add_heading("Attendance 2026-09-03", 1)
    t = d.add_table(rows=1, cols=5); t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, ["Employee ID", "Name", "Department", "Date", "Status"]): c.text = h
    for r in rows("2026-09-03"):
        for c, v in zip(t.add_row().cells, r): c.text = v
    d.save(out / "attendance.docx")
    pdf = fitz.open(); p = pdf.new_page(); p.insert_text((50, 72), "Attendance sheet 2026-09-04", fontsize=14)
    for i, r in enumerate(rows("2026-09-04")): p.insert_text((50, 100 + 20 * i), " | ".join(r), fontsize=11)
    pdf.save(out / "attendance_text.pdf")
    png, size = scan_png("2026-09-05"); (out / "attendance_scan.png").write_bytes(png)
    png, size = scan_png("2026-09-08"); pdf = fitz.open()
    p = pdf.new_page(width=size[0] * .6, height=size[1] * .6); p.insert_image(p.rect, stream=png); pdf.save(out / "attendance_scan.pdf")

    df = pd.DataFrame([r for dd in DAYS for r in rows(dd)], columns=COLS); pres = df.status.eq("Present")
    top = lambda col: (lambda g: sorted(g[g == g.max()].index))(df.assign(p=pres).groupby(col).p.mean().mul(100).round(1))
    exp = {"files": {"attendance.csv": "2026-09-01", "attendance.xlsx": "2026-09-02", "attendance.docx": "2026-09-03",
                     "attendance_text.pdf": "2026-09-04", "attendance_scan.png": "2026-09-05 (OCR)", "attendance_scan.pdf": "2026-09-08 (OCR)"},
           "total_records": len(df),
           "Q: Who was present on 2026-09-03?": sorted(df[(df.date == "2026-09-03") & pres].employee_name),
           "Q: Average attendance 2026-09-01 to 2026-09-08 (%)": round(pres.mean() * 100, 1),
           "Q: Average attendance 2026-09-01 to 2026-09-03 (%)": round(pres[df.date <= "2026-09-03"].mean() * 100, 1),
           "Q: Which employee had the highest attendance?": top("employee_name"),
           "Q: Which department had the highest attendance?": top("department"),
           "Q: Who was present on 2026-12-25?": "NOT FOUND - answer must say the data does not contain it, with no sources"}
    (out / "expected_results.json").write_text(json.dumps(exp, indent=2))
    (out / "canonical_schema.json").write_text(json.dumps({
        "record": {"employee_id": "string|''", "employee_name": "string (required, case-insensitive identity)", "department": "string|''",
                   "date": "YYYY-MM-DD", "status": "Present|Absent|Leave", "file": "source file", "ref": "sheet/row, table/row, page/lines or OCR lines"},
        "key": "(user, lower(employee_name), date) - re-uploads overwrite", "attendance_pct": "Present / (Present+Absent+Leave) * 100"}, indent=2))
    print("Samples written to", out)

if __name__ == "__main__":
    main(*sys.argv[1:2])
