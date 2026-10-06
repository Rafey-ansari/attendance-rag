"""Core: SQLite storage, file parsing + OCR, Chroma vector store, LangGraph QA pipeline."""
import base64, io, json, os, re, sqlite3, uuid
from datetime import date
from pathlib import Path
from typing import TypedDict
from openai import OpenAI

import pandas as pd
from dotenv import load_dotenv
from langgraph.graph import END, StateGraph

load_dotenv('.env', override=True)
DATA = Path(os.getenv("DATA_DIR", "data"))
(DATA / "uploads").mkdir(parents=True, exist_ok=True)
MODEL = os.getenv("OPENAI_MODEL", "gpt-4o-mini")
EMBED_MODEL = os.getenv("EMBED_MODEL", "text-embedding-3-small")
VISION = os.getenv("VISION_FALLBACK", "true").lower() == "true"
ALLOWED = {".csv", ".tsv", ".txt", ".xlsx", ".docx", ".pdf", ".png", ".jpg", ".jpeg"}
NOT_FOUND = "The supplied data does not contain the answer to this question."

# ---------------------------------------------------------------- database
def db():
    c = sqlite3.connect(DATA / "app.db")
    c.row_factory = sqlite3.Row
    return c

def q(sql, args=()):
    """Run one statement. SELECT -> list of dicts, write -> lastrowid. Always parameterised."""
    c = db()
    try:
        cur = c.execute(sql, args)
        c.commit()
        return [dict(r) for r in cur.fetchall()] if cur.description else cur.lastrowid
    finally:
        c.close()

def init_db():
    c = db()
    c.executescript("""
    CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY AUTOINCREMENT, email TEXT UNIQUE, pw TEXT, created TEXT DEFAULT CURRENT_TIMESTAMP);
    CREATE TABLE IF NOT EXISTS files(id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INT, name TEXT, path TEXT, n_records INT, created TEXT DEFAULT CURRENT_TIMESTAMP);
    CREATE TABLE IF NOT EXISTS records(user_id INT, emp_key TEXT, employee_id TEXT, employee_name TEXT, department TEXT,
        date TEXT, status TEXT, file TEXT, ref TEXT, PRIMARY KEY(user_id, emp_key, date));
    CREATE TABLE IF NOT EXISTS messages(id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INT, question TEXT, answer TEXT, sources TEXT, found INT, created TEXT DEFAULT CURRENT_TIMESTAMP);
    CREATE TABLE IF NOT EXISTS feedback(id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INT, message_id INT, rating INT, comment TEXT, created TEXT DEFAULT CURRENT_TIMESTAMP);
    """)
    c.commit(); c.close()

# ---------------------------------------------------------------- OpenAI + vector store
_client = _chroma = _ocr = None

def client():
    global _client

    if _client is None:
        api_key = os.getenv("OPENAI_API_KEY")

        if not api_key:
            raise RuntimeError("OPENAI_API_KEY was not loaded")

        print("Using OpenAI key:", api_key[:8] + "..." + api_key[-4:])
        _client = OpenAI(api_key=api_key)
    return _client

def llm(system, user, json_mode=True):
    kw = {"response_format": {"type": "json_object"}} if json_mode else {}
    r = client().chat.completions.create(model=MODEL, temperature=0, **kw,
        messages=[{"role": "system", "content": system}, {"role": "user", "content": user}])
    t = r.choices[0].message.content
    return json.loads(t) if json_mode else t

def embed(texts):
    out = []
    for i in range(0, len(texts), 96):
        out += [d.embedding for d in client().embeddings.create(model=EMBED_MODEL, input=texts[i:i + 96]).data]
    return out

def coll(uid):
    global _chroma
    if _chroma is None:
        import chromadb
        _chroma = chromadb.PersistentClient(path=str(DATA / "chroma"))
    return _chroma.get_or_create_collection(f"user_{uid}", metadata={"hnsw:space": "cosine"})

# ---------------------------------------------------------------- OCR (RapidOCR, open source, ONNX)
def ocr_image(img_bytes):
    """Return (text, avg_confidence). Words on the same visual line are joined, so table rows stay rows."""
    global _ocr
    import numpy as np
    from PIL import Image
    from rapidocr_onnxruntime import RapidOCR
    _ocr = _ocr or RapidOCR()
    res, _ = _ocr(np.array(Image.open(io.BytesIO(img_bytes)).convert("RGB")))
    if not res:
        return "", 0.0
    items = sorted(((min(p[1] for p in b), max(p[1] for p in b), min(p[0] for p in b), t, float(s)) for b, t, s in res))
    lines = []
    for top, bot, left, t, _s in items:
        if lines and top - lines[-1][0] < 0.6 * (bot - top):
            lines[-1][1].append((left, t))
        else:
            lines.append((top, [(left, t)]))
    text = "\n".join("  ".join(t for _, t in sorted(w)) for _, w in lines)
    return text, sum(i[4] for i in items) / len(items)

def read_image(img_bytes):
    text, conf = ocr_image(img_bytes)
    if conf < 0.75 and VISION and os.getenv("OPENAI_API_KEY"):  # handwriting: optional vision fallback
        from PIL import Image
        buf = io.BytesIO(); Image.open(io.BytesIO(img_bytes)).convert("RGB").save(buf, "PNG")
        b64 = base64.b64encode(buf.getvalue()).decode()
        r = client().chat.completions.create(model=MODEL, messages=[{"role": "user", "content": [
            {"type": "text", "text": "Transcribe all text in this attendance sheet row by row, exactly as written. Tables as pipe-separated rows."},
            {"type": "image_url", "image_url": {"url": "data:image/png;base64," + b64}}]}])
        text = r.choices[0].message.content or text
    return text

# ---------------------------------------------------------------- parsing -> canonical records
SYN = {"employee_id": ["employee id", "emp id", "emp code", "empid", "employee code", "id"],
       "employee_name": ["employee name", "employee", "name", "emp name", "staff"],
       "department": ["department", "dept", "team", "division"],
       "date": ["date", "attendance date", "day"],
       "status": ["status", "attendance", "present/absent", "presence"]}
PRESENT, ABSENT, LEAVE = {"p", "present", "1", "yes", "y", "wfh", "late"}, {"a", "absent", "0", "no", "n"}, {"l", "leave", "pl", "sl", "cl", "on leave"}

def norm_status(s):
    s = str(s or "").strip().lower()
    return "Present" if s in PRESENT else "Absent" if s in ABSENT else "Leave" if s in LEAVE else None

def to_iso(v):
    s = str(v or "").strip()
    m = re.match(r"\d{4}-\d{2}-\d{2}", s)
    if m:
        return m.group(0)
    d = pd.to_datetime(s, dayfirst=True, errors="coerce")  # ambiguous dates are read day-first
    return None if pd.isna(d) else d.date().isoformat()

def map_cols(cols):
    low = {c: re.sub(r"[_\s]+", " ", str(c).strip().lower()) for c in cols}
    m = {}
    for k, names in SYN.items():
        for c, l in low.items():
            if l in names and c not in m.values():
                m[k] = c; break
    return m

def table_records(df, label):
    """Deterministic mapping of a table with recognisable headers. Returns [] if headers are not recognised."""
    m = map_cols(df.columns)
    if "date" not in m or "status" not in m or not ({"employee_name", "employee_id"} & m.keys()):
        return []
    out = []
    for pos, (_, r) in enumerate(df.iterrows()):
        d, s = to_iso(r[m["date"]]), norm_status(r[m["status"]])
        eid = str(r[m["employee_id"]]).strip() if "employee_id" in m else ""
        name = str(r[m["employee_name"]]).strip() if "employee_name" in m else eid
        if d and s and name:
            out.append(dict(employee_id=eid, employee_name=name, department=str(r[m["department"]]).strip() if "department" in m else "",
                            date=d, status=s, ref=f"{label}, row {pos + 2}"))
    return out

def chunk_lines(lines, label, size=30):
    lines = [l for l in lines if l.strip()]
    return [(f"{label}, lines {i + 1}-{min(i + size, len(lines))}", "\n".join(lines[i:i + size])) for i in range(0, len(lines), size)]

def df_segments(df, label, size=40):
    lines = [f"row {i + 2}: " + "; ".join(f"{c}={v}" for c, v in r.items() if v) for i, (_, r) in enumerate(df.iterrows())]
    return [(f"{label}, rows {a + 2}-{min(a + size, len(lines)) + 1}", "\n".join(lines[a:a + size])) for a in range(0, len(lines), size)]

def parse(path, name):
    """-> (segments [(ref, text)], records). Records are non-empty only when headers were recognised."""
    ext, segs, recs = Path(name).suffix.lower(), [], []
    if ext in (".csv", ".tsv"):
        df = pd.read_csv(path, dtype=str, sep=None, engine="python").fillna("")
        recs, segs = table_records(df, name), df_segments(df, name)
    elif ext == ".txt":
        segs = chunk_lines(Path(path).read_text(errors="ignore").splitlines(), name)
    elif ext == ".xlsx":
        for sh, df in pd.read_excel(path, sheet_name=None, dtype=str).items():
            df = df.fillna("")
            recs += table_records(df, f"{name} › {sh}"); segs += df_segments(df, f"{name} › {sh}")
    elif ext == ".docx":
        import docx
        d = docx.Document(path)
        segs += chunk_lines([p.text for p in d.paragraphs], f"{name} › text")
        for n, t in enumerate(d.tables, 1):
            rows = [[c.text.strip() for c in r.cells] for r in t.rows]
            if len(rows) > 1:
                df = pd.DataFrame(rows[1:], columns=rows[0])
                recs += table_records(df, f"{name} › table {n}"); segs += df_segments(df, f"{name} › table {n}")
    elif ext == ".pdf":
        import fitz
        for pn, page in enumerate(fitz.open(path), 1):
            t, label = page.get_text().strip(), f"{name}, page {pn}"
            if len(t) < 20:  # scanned page -> render + OCR
                t, label = read_image(page.get_pixmap(dpi=200).tobytes("png")), label + " (OCR)"
            segs += chunk_lines(t.splitlines(), label)
    else:  # images
        segs = chunk_lines(read_image(Path(path).read_bytes()).splitlines(), f"{name} (OCR)")
    return segs, recs

EXTRACT_SYS = ("You extract attendance records from messy text (table rows, OCR output, free text). Return JSON: "
    '{"records":[{"employee_id":string|null,"employee_name":string,"department":string|null,"date":"YYYY-MM-DD","status":"Present|Absent|Leave"}]}. '
    "Ambiguous dates are day-first. If a header gives the date for the whole sheet, apply it to every row. "
    "Expand grids (employees x dates) into one record per cell. Never invent rows; skip anything unreadable.")

def ingest(uid, path, name):
    segs, recs = parse(path, name)
    if not recs:  # unknown layout / text / OCR -> LLM normalisation to the canonical schema
        for ref, text in segs:
            for r in llm(EXTRACT_SYS, f"SOURCE: {ref}\n{text}").get("records", []):
                d, s = to_iso(r.get("date")), norm_status(r.get("status"))
                nm = (r.get("employee_name") or r.get("employee_id") or "").strip()
                if d and s and nm:
                    recs.append(dict(employee_id=r.get("employee_id") or "", employee_name=nm,
                                     department=r.get("department") or "", date=d, status=s, ref=ref))
    for old in q("SELECT id FROM files WHERE user_id=? AND name=?", (uid, name)):
        q("DELETE FROM files WHERE id=?", (old["id"],))
    c = db()
    c.executemany("INSERT OR REPLACE INTO records VALUES (?,?,?,?,?,?,?,?,?)",
        [(uid, r["employee_name"].lower(), r["employee_id"], r["employee_name"], r["department"], r["date"], r["status"], name, r["ref"]) for r in recs])
    c.commit(); c.close()
    q("INSERT INTO files(user_id,name,path,n_records) VALUES (?,?,?,?)", (uid, name, str(path), len(recs)))
    col = coll(uid)
    col.delete(where={"file": name})
    if segs:
        col.add(ids=[uuid.uuid4().hex for _ in segs], documents=[t for _, t in segs],
                metadatas=[{"file": name, "ref": r} for r, _ in segs], embeddings=embed([t for _, t in segs]))
    return len(recs)

def delete_file(uid, fid):
    f = q("SELECT * FROM files WHERE id=? AND user_id=?", (fid, uid))
    if not f:
        return False
    q("DELETE FROM records WHERE user_id=? AND file=?", (uid, f[0]["name"]))
    coll(uid).delete(where={"file": f[0]["name"]})
    q("DELETE FROM files WHERE id=?", (fid,))
    Path(f[0]["path"]).unlink(missing_ok=True)
    return True

# ---------------------------------------------------------------- deterministic attendance maths
def compute(uid, p):
    """Exact answers from canonical records. attendance % = Present / (Present+Absent+Leave). Returns None if no data."""
    df = pd.DataFrame(q("SELECT * FROM records WHERE user_id=?", (uid,)))
    if df.empty:
        return None
    op = p.get("operation")
    if p.get("department"):
        df = df[df.department.str.lower() == str(p["department"]).lower()]
    if p.get("employee"):
        e = str(p["employee"]).lower()
        df = df[df.employee_name.str.lower().str.contains(e, regex=False) | (df.employee_id.str.lower() == e)]
    if p.get("start"): df = df[df.date >= p["start"]]
    if p.get("end"): df = df[df.date <= p["end"]]
    if df.empty:
        return None
    pct = lambda s: round((s == "Present").mean() * 100, 1)
    if op in ("present_on_date", "absent_on_date", "employee_status") and p.get("date"):
        day = df[df.date == p["date"]]
        if day.empty:
            return None
        if op == "employee_status":
            return {"summary": "; ".join(f"{r.employee_name}: {r.status} on {r.date}" for r in day.itertuples()), "rows": day}
        want = "Present" if op == "present_on_date" else "Absent"
        names = sorted(day[day.status == want].employee_name)
        return {"summary": f"{len(names)} of {len(day)} employees recorded were {want} on {p['date']}: {', '.join(names) or 'none'}", "rows": day}
    if op == "avg_attendance":
        return {"summary": f"Average attendance {pct(df.status)}% ({(df.status == 'Present').sum()} Present of {len(df)} records, {df.date.min()} to {df.date.max()})", "rows": df}
    if op in ("top_employee", "top_department"):
        col = "employee_name" if op == "top_employee" else "department"
        d2 = df[df[col] != ""]
        if d2.empty:
            return None
        g = d2.groupby(col).status.apply(pct).sort_values(ascending=False)
        win = g[g == g.max()]
        return {"summary": f"Highest {col.split('_')[0]} attendance: {', '.join(win.index)} at {g.max()}%. Ranking: " +
                ", ".join(f"{k} {v}%" for k, v in g.head(5).items()), "rows": d2[d2[col].isin(win.index)]}
    return None

# ---------------------------------------------------------------- LangGraph pipeline
class S(TypedDict, total=False):
    uid: int; question: str; plan: dict; summary: str; rows: list; chunks: list; answer: str; found: bool; sources: list

PLAN_SYS = ("You are a query planner for an attendance database. Today is {today}. Return JSON: "
    '{{"operation":"present_on_date|absent_on_date|avg_attendance|top_employee|top_department|employee_status|other",'
    '"date":"YYYY-MM-DD"|null,"start":"YYYY-MM-DD"|null,"end":"YYYY-MM-DD"|null,"employee":string|null,"department":string|null}}. '
    'Use "other" for anything that is not one of these attendance questions.')

def n_plan(s):
    return {
        "plan": llm(
            PLAN_SYS + feedback_prompt(s["uid"]),
            s["question"]
        )
    }

def n_compute(s):
    r = compute(s["uid"], s["plan"])
    return {"summary": r["summary"], "rows": r["rows"].to_dict("records")} if r else {}

def n_retrieve(s):
    col = coll(s["uid"])
    if col.count() == 0:
        return {"chunks": []}
    r = col.query(query_embeddings=embed([s["question"]]), n_results=min(5, col.count()))
    return {"chunks": [{"file": m["file"], "ref": m["ref"], "text": d} for d, m in zip(r["documents"][0], r["metadatas"][0])]}

def feedback_prompt(uid):
    rows = q(
        """
        SELECT 
            m.question,
            m.answer,
            f.rating,
            f.comment
        FROM feedback f
        JOIN messages m 
            ON m.id = f.message_id
        WHERE f.user_id=?
        ORDER BY f.id DESC
        LIMIT 10
        """,
        (uid,)
    )

    if not rows:
        return "(No previous feedback from this user.)"

    feedback = []

    for r in rows:
        rating = "POSITIVE" if r["rating"] == 1 else "NEGATIVE"

        comment = r["comment"].strip() if r["comment"] else "(No written comment)"

        feedback.append(
            f"""
Previous question:
{r["question"]}

Previous answer:
{r["answer"]}

User feedback: {rating}
User comment: {comment}
""".strip()
        )

    return "\n\n---\n\n".join(feedback)

ANSWER_SYS = (
    "You are an attendance assistant. "
    "Answer ONLY from the CONTEXT. "
    "Attendance % = Present/(Present+Absent+Leave). "

    "If a COMPUTED RESULT is present, trust it exactly. "
    "If the CONTEXT does not support an answer, say the supplied data does not "
    "contain it and set found=false. "
    "Never guess. "

    "Use the USER FEEDBACK section to improve the style, structure, clarity, "
    "and level of detail of your answer. "
    "Do NOT treat user feedback as factual attendance data. "
    "Feedback only tells you how the user prefers answers to be presented. "

    'Return JSON {"answer":string,"found":boolean}.\n\n'

    "USER FEEDBACK TO RESPECT:\n"
)

def n_answer(s):
    ctx = []

    if s.get("summary"):
        ctx.append("COMPUTED RESULT:\n" + s["summary"])

        ctx.append(
            "EVIDENCE ROWS:\n" +
            "\n".join(
                f"- {r['employee_name']} | "
                f"{r['department']} | "
                f"{r['date']} | "
                f"{r['status']} | "
                f"{r['file']} :: {r['ref']}"
                for r in s["rows"][:40]
            )
        )

    ctx += [
        f"DOCUMENT EXCERPT [{c['file']} :: {c['ref']}]:\n{c['text']}"
        for c in s.get("chunks", [])
    ]

    if not ctx:
        return {
            "answer": NOT_FOUND,
            "found": False,
            "sources": []
        }

    feedback = feedback_prompt(s["uid"])

    prompt = (
        "CONTEXT:\n"
        + "\n\n".join(ctx)
        + "\n\n"
        + "USER FEEDBACK FROM PREVIOUS INTERACTIONS:\n"
        + feedback
        + "\n\n"
        + "CURRENT QUESTION:\n"
        + s["question"]
    )

    out = llm(ANSWER_SYS, prompt)

    found = bool(out.get("found"))

    src = []

    if found:
        if s.get("summary"):
            seen = {}

            for r in s["rows"]:
                seen.setdefault(
                    (r["file"], r["ref"]),
                    f"{r['employee_name']}, {r['date']}, {r['status']}"
                )

            src = [
                {
                    "file": f,
                    "ref": ref,
                    "text": t
                }
                for (f, ref), t in list(seen.items())[:20]
            ]

        else:
            src = [
                {
                    "file": c["file"],
                    "ref": c["ref"],
                    "text": c["text"][:200]
                }
                for c in s["chunks"]
            ]

    return {
        "answer": out.get("answer") or NOT_FOUND,
        "found": found,
        "sources": src
    }

def _build():
    g = StateGraph(S)
    for n, f in [("plan", n_plan), ("compute", n_compute), ("retrieve", n_retrieve), ("answer", n_answer)]:
        g.add_node(n, f)
    g.set_entry_point("plan")
    g.add_edge("plan", "compute")
    g.add_conditional_edges("compute", lambda s: "answer" if s.get("summary") else "retrieve", {"answer": "answer", "retrieve": "retrieve"})
    g.add_edge("retrieve", "answer")
    g.add_edge("answer", END)
    return g.compile()

GRAPH = _build()

def ask(uid, question):
    if not q("SELECT 1 FROM files WHERE user_id=?", (uid,)):
        out = {"answer": "Upload at least one attendance file first.", "found": False, "sources": []}
    else:
        out = GRAPH.invoke({"uid": uid, "question": question})
    mid = q("INSERT INTO messages(user_id,question,answer,sources,found) VALUES (?,?,?,?,?)",
            (uid, question, out["answer"], json.dumps(out["sources"]), int(out["found"])))
    return {"id": mid, "answer": out["answer"], "found": out["found"], "sources": out["sources"]}
